"""Version-aware (v1/v2) Excel workbook import, alongside Tempo's existing CSV pipeline — not instead
of it (brief: "retain CSV support without content-based row deletion").

Arch's reviewed templates are multi-sheet workbooks where customers are told to upload the WHOLE
file, not one sheet at a time, with a fixed structural layout per sheet: row 1 is the column header,
rows 2-4 are guidance/example content (not data — a customer who leaves the example row untouched
must not get it imported), and real data begins at row 5. This module was built from that literal,
stated layout; it was not checked against an actual Arch source file, which this programme was never
given access to (the original brief explicitly said not to block on that access — see
docs/order-driven-planning.md). The skip is purely POSITIONAL (row 2, 3 and 4, always, regardless of
what they contain) — never content-based, the same principle the CSV pipeline already follows (see
the "preserve ignored example rows" finding in docs/order-driven-planning.md: a content-based guess
risks discarding genuine customer data).

Each recognised sheet is converted into the exact same (headers, rows) shape `read_csv` produces, then
fed through the SAME staging pipeline (`app.imports.engine.stage`) as any CSV upload — one
`ImportBatch` per sheet, previewed and applied through the existing endpoints, unchanged. "Version-
aware v1/v2" is satisfied the same way the CSV contracts already are: every field this programme added
(function/flow/required_skill, award, day_rates, award_rules, award_eligibility_restrictions, ...) is
optional, so a v1 sheet missing those columns and a v2 sheet with them both validate against the one
contract definition.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.imports.parse import ImportProblem

# Row 1 = header; rows 2-4 = guidance/example content, always skipped by position; data from row 5.
HEADER_ROW = 1
DATA_START_ROW = 5

# Sheet tab name (case/space-insensitive) -> (data_class, entity). Covers every CONTRACTS entry a
# workbook upload could plausibly carry; a handful of the more likely v1/v2 naming variants are
# included as synonyms, the same spirit as each Field's own `synonyms` for column headers.
_SHEET_ALIASES: dict[str, tuple[str, str]] = {
    "sites": ("master", "sites"), "warehouses": ("master", "sites"),
    "customers": ("master", "customers"),
    "zones": ("master", "zones"), "areas": ("master", "zones"),
    "work standards": ("master", "work_standards"), "work_standards": ("master", "work_standards"), "activities": ("master", "work_standards"), "activity standards": ("master", "work_standards"),
    "workers": ("master", "workers"), "staff": ("master", "workers"), "employees": ("master", "workers"),
    "rates": ("master", "rates"), "pay rates": ("master", "rates"), "labour rates": ("master", "rates"),
    "availability": ("master", "availability"), "leave": ("master", "availability"),
    "weekly availability": ("master", "weekly_availability"),
    "operating calendar": ("master", "operating_calendar"), "operating hours": ("master", "operating_calendar"),
    "shift templates": ("master", "shift_templates"), "shifts": ("master", "shift_templates"),
    "shift breaks": ("master", "shift_breaks"), "breaks": ("master", "shift_breaks"),
    "equipment": ("master", "equipment"),
    "activity roles": ("master", "activity_roles"),
    "process templates": ("master", "process_templates"), "processes": ("master", "process_templates"),
    "process steps": ("master", "process_steps"), "tasks": ("master", "process_steps"),
    "orders": ("master", "orders"), "outbound orders": ("master", "orders"),
    "worker activity rates": ("master", "worker_activity_rates"), "task rates": ("master", "worker_activity_rates"), "individual rates": ("master", "worker_activity_rates"),
    "unit conversions": ("master", "unit_conversions"),
    "fill priorities": ("master", "fill_priorities"),
    "absenteeism": ("master", "absenteeism"),
    "headcount limits": ("master", "headcount_limits"),
    "grade rates": ("master", "grade_rates"),
    "productivity loss": ("master", "productivity_loss"),
    "staging capacity": ("master", "staging_capacity"),
    "staging movements": ("master", "staging_movements"),
    "indirect headcount": ("master", "indirect_headcount"), "supervision": ("master", "indirect_headcount"),
    "day rates": ("master", "day_rates"),
    "award rules": ("master", "award_rules"), "awards": ("master", "award_rules"),
    "award eligibility restrictions": ("master", "award_eligibility_restrictions"),
}

# Dependency order for staging multiple sheets from ONE workbook in a single pass — NOT the same as
# contracts.DATA_CLASSES's tuple, which is grouped for display, not safe processing order (e.g. it
# lists "workers" before "sites", but a worker's home_site must already exist). A sheet not in this
# list (there should be none) sorts after everything named here, in whatever order it was found.
_PROCESSING_ORDER: tuple[str, ...] = (
    "sites", "customers", "zones", "work_standards", "workers", "rates", "availability", "weekly_availability",
    "operating_calendar", "shift_templates", "shift_breaks", "equipment", "activity_roles",
    "process_templates", "process_steps", "orders", "worker_activity_rates", "unit_conversions",
    "fill_priorities", "absenteeism", "headcount_limits", "grade_rates", "productivity_loss",
    "staging_capacity", "staging_movements", "indirect_headcount", "day_rates", "award_rules", "award_eligibility_restrictions",
)


def sheet_entity(sheet_name: str) -> tuple[str, str] | None:
    """The (data_class, entity) a sheet tab name maps to, or None if unrecognised — an unrecognised
    sheet is reported back to the caller, never silently dropped (see xlsx_stage/xlsx_inspect)."""
    return _SHEET_ALIASES.get(sheet_name.strip().lower())


def processing_rank(entity: str) -> int:
    return _PROCESSING_ORDER.index(entity) if entity in _PROCESSING_ORDER else len(_PROCESSING_ORDER)


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_xlsx(data: bytes) -> dict[str, tuple[list[str], list[dict[str, str]]]]:
    """Every sheet in the workbook, each as (headers, rows) — the same shape `read_csv` returns, so
    every downstream caller (mapping suggestion, `engine.stage`) is unchanged. Row 1 is the header;
    rows 2-4 are skipped structurally regardless of content; data starts at row 5. A genuinely blank
    data row (every cell empty) is skipped the same way `read_csv` skips one."""
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several distinct exception types for a bad file
        raise ImportProblem("the file is not a readable .xlsx workbook") from exc
    sheets: dict[str, tuple[list[str], list[dict[str, str]]]] = {}
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            sheets[ws.title] = ([], [])
            continue
        headers = [_cell_to_str(h) for h in header_cells]
        rows: list[dict[str, str]] = []
        for row_no, cells in enumerate(rows_iter, start=2):
            if row_no < DATA_START_ROW:
                continue  # rows 2-4: guidance/example content, skipped by position
            values = [_cell_to_str(c) for c in cells]
            if not any(values):
                continue  # a genuinely blank data row — same convention as read_csv
            rows.append({h: (values[i] if i < len(values) else "") for i, h in enumerate(headers) if h})
        sheets[ws.title] = (headers, rows)
    wb.close()
    return sheets
